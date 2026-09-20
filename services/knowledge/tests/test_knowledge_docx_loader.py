import io
import zipfile
from collections.abc import Callable

import pytest
from docx import Document as WordDocument

from factotum.contracts.errors import DocumentTooLargeError, UnsupportedDocumentError
from factotum.knowledge.document import Limits
from factotum.knowledge.loaders.docx import parse_docx

DocxFactory = Callable[..., bytes]


def test_paragraphs_headings_and_the_title_are_read(docx: DocxFactory) -> None:
    parsed = parse_docx(docx(), Limits())

    assert parsed.title == "Router guide"
    assert [(b.text, b.level) for b in parsed.blocks if b.kind == "heading"] == [
        ("Resetting", 1),
        ("Factory reset", 2),
    ]
    paragraph = next(b for b in parsed.blocks if b.text.startswith("Hold"))
    assert paragraph.headings == ("Resetting", "Factory reset")


def test_without_a_document_title_it_is_the_first_heading(docx: DocxFactory) -> None:
    assert parse_docx(docx(title=None), Limits()).title == "Resetting"


def test_without_either_there_is_no_title(docx: DocxFactory) -> None:
    assert parse_docx(docx(title=None, headings=False), Limits()).title is None


def test_a_list_keeps_its_bullets(docx: DocxFactory) -> None:
    parsed = parse_docx(docx(bullets=["first", "second"]), Limits())

    assert "- first" in [b.text for b in parsed.blocks]


def test_a_table_is_a_block_of_rows(docx: DocxFactory) -> None:
    parsed = parse_docx(docx(table=[["Model", "Button"], ["A1", "left"]]), Limits())

    table = next(b for b in parsed.blocks if b.kind == "table")
    assert table.text == "Model | Button\nA1 | left"


def test_the_title_style_counts_as_a_heading_of_the_first_level() -> None:
    document = WordDocument()
    document.add_heading("The manual", level=0)
    document.add_paragraph("Some text.")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_docx(buffer.getvalue(), Limits())

    assert [(b.text, b.level) for b in parsed.blocks if b.kind == "heading"] == [("The manual", 1)]


def test_a_table_with_nothing_in_it_is_dropped(docx: DocxFactory) -> None:
    parsed = parse_docx(docx(table=[["", " "], ["", ""]]), Limits())

    assert not any(b.kind == "table" for b in parsed.blocks)


def test_empty_paragraphs_are_dropped(docx: DocxFactory) -> None:
    parsed = parse_docx(docx(paragraphs=["", "   ", "Only this"]), Limits())

    assert [b.text for b in parsed.blocks if b.kind == "text"] == ["Only this"]


def test_something_that_is_not_a_word_document_is_refused() -> None:
    not_word = io.BytesIO()
    with zipfile.ZipFile(not_word, "w") as archive:
        archive.writestr("word/document.xml", "this is not xml")

    with pytest.raises(UnsupportedDocumentError, match="could not be read"):
        parse_docx(not_word.getvalue(), Limits())


def test_a_damaged_archive_is_refused() -> None:
    with pytest.raises(UnsupportedDocumentError, match="damaged"):
        parse_docx(b"PK\x03\x04 not really", Limits())


def bomb(size: int) -> bytes:
    """A .docx whose one big part is nothing but zeros: tiny to send, huge to open."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * size)
    return buffer.getvalue()


def test_an_archive_that_expands_to_too_much_is_refused_before_it_is_opened() -> None:
    with pytest.raises(DocumentTooLargeError, match="expands to too much"):
        parse_docx(bomb(20_000), Limits(max_zip_bytes=10_000))


def test_an_archive_that_compresses_suspiciously_well_is_refused() -> None:
    with pytest.raises(DocumentTooLargeError, match="suspiciously well"):
        parse_docx(bomb(5_000_000), Limits())


def test_an_archive_with_too_many_parts_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<w/>")
        for number in range(20_001):
            archive.writestr(f"part{number}.xml", "")

    with pytest.raises(DocumentTooLargeError, match="too many parts"):
        parse_docx(buffer.getvalue(), Limits())
