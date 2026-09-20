"""Documents built for the tests, so that no binary fixture lives in the repository."""

import io
import textwrap
from collections.abc import Callable, Mapping, Sequence

import pytest
from docx import Document as WordDocument
from reportlab.lib.pagesizes import A4
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.pdfgen import canvas

PdfFactory = Callable[..., bytes]
DocxFactory = Callable[..., bytes]


def build_pdf(
    chapters: Mapping[str, Sequence[str]],
    *,
    header: str | None = "The Book",
    numbered: bool = True,
    password: str | None = None,
    contents: Sequence[str] = (),
    sections: bool = False,
) -> bytes:
    """A PDF with a chapter per page, an outline, and the running header and page numbers a printed
    book has. ``contents`` becomes a first page of table-of-contents lines with dot leaders."""
    buffer = io.BytesIO()
    encryption = StandardEncryption(password, ownerPassword="owner") if password else None
    pdf = canvas.Canvas(buffer, pagesize=A4, encrypt=encryption)
    page = 0

    if contents:
        page += 1
        y = 800
        for line in contents:
            pdf.drawString(72, y, line)
            y -= 14
        pdf.showPage()

    for title, paragraphs in chapters.items():
        page += 1
        key = f"page{page}"
        pdf.bookmarkPage(key)
        pdf.addOutlineEntry(title, key, level=0)
        if sections:
            pdf.bookmarkPage(f"{key}-about")  # reportlab keeps one title per destination
            pdf.addOutlineEntry(f"About {title}", f"{key}-about", level=1)
        y = 800
        if header:
            pdf.drawString(72, y, f"{header}: {title}")
            y -= 20
        pdf.setFont("Helvetica", 11)
        for paragraph in paragraphs:
            for line in textwrap.wrap(paragraph, 95) or [""]:
                pdf.drawString(72, y, line)
                y -= 14
            y -= 6
        if numbered:
            pdf.drawString(300, 30, str(page))
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def build_blank_pdf(pages: int = 2) -> bytes:
    """Pages with nothing in them, like the scan of a page that has not been through OCR."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for _ in range(pages):
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def build_docx(
    *,
    title: str | None = "Router guide",
    headings: bool = True,
    table: Sequence[Sequence[str]] = (),
    bullets: Sequence[str] = (),
    paragraphs: Sequence[str] = ("Hold the reset button for ten seconds.",),
) -> bytes:
    document = WordDocument()
    if title:
        document.core_properties.title = title
    if headings:
        document.add_heading("Resetting", level=1)
        document.add_heading("Factory reset", level=2)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for bullet in bullets:
        document.add_paragraph(bullet, style="List Bullet")
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column, cell in enumerate(row):
                grid.cell(row_index, column).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def pdf() -> PdfFactory:
    return build_pdf


@pytest.fixture
def docx() -> DocxFactory:
    return build_docx


@pytest.fixture
def blank_pdf() -> PdfFactory:
    return build_blank_pdf
