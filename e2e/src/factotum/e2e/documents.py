"""Documents as a person has them: a PDF and a Word file, made the way a program that writes them
does, so that the knowledge service reads them with the loaders it has (in a process of its own)."""

import io
from collections.abc import Sequence

from docx import Document as WordDocument
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

PDF = "application/pdf"
WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def a_pdf(pages: Sequence[tuple[str, Sequence[str]]]) -> bytes:
    """A PDF with a section on each page: its title (in the outline, as in a book, and at the top of
    the page) and its paragraphs, each a line apart from the next."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for number, (title, paragraphs) in enumerate(pages, start=1):
        key = f"page-{number}"
        pdf.bookmarkPage(key)
        pdf.addOutlineEntry(title, key)
        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(72, 780, title)
        pdf.setFont("Helvetica", 11)
        height = 740
        for paragraph in paragraphs:
            pdf.drawString(72, height, paragraph)
            height -= 34
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def a_word_document(sections: Sequence[tuple[str, Sequence[str]]]) -> bytes:
    """A Word file with a heading and some paragraphs for each section."""
    document = WordDocument()
    for title, paragraphs in sections:
        document.add_heading(title, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
